from flask import Flask, render_template, request, redirect, url_for, session, flash, send_file
import io
import sqlite3
from datetime import datetime
import pandas as pd
import os
import string
import secrets
from functools import wraps

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
DB_YOLU = os.path.join(BASE_DIR, 'meb_otomasyon.db')

app = Flask(__name__)
app.secret_key = 'meb_super_gizli_anahtar_2027' 

def guvenli_sifre_olustur(uzunluk=8):
    harfler = string.ascii_letters
    rakamlar = string.digits
    ozel_karakterler = "!@#$%^&*"
    tum_karakterler = harfler + rakamlar + ozel_karakterler

    while True:
        sifre = ''.join(secrets.choice(tum_karakterler) for _ in range(uzunluk))
        if (any(c.islower() for c in sifre) and 
            any(c.isupper() for c in sifre) and 
            any(c.isdigit() for c in sifre) and 
            any(c in ozel_karakterler for c in sifre)):
            return sifre

KISITLI_ROLLER = ["Şef", "Memur"]

def tam_yetki_gerekli(f):
    @wraps(f)
    def decorated_function(*args, **kwargs):
        kullanici_rolu = session.get('kullanici_rolu')
        if kullanici_rolu in KISITLI_ROLLER:
            flash("HATA: Bu işlem (Silme/Onaylama) için yetkiniz bulunmamaktadır!", "danger")
            return redirect(url_for('ana_panel'))
        return f(*args, **kwargs)
    return decorated_function

def veritabani_kurulum():
    conn = sqlite3.connect(DB_YOLU)
    cursor = conn.cursor()
    
    cursor.execute('''
        CREATE TABLE IF NOT EXISTS tbl_kullanicilar (
            kullanici_adi VARCHAR(50) PRIMARY KEY,
            sifre VARCHAR(50),
            rol VARCHAR(50),
            ad_soyad VARCHAR(100)
        )
    ''')
    
    try:
        cursor.execute("ALTER TABLE tbl_nakil_talepleri ADD COLUMN yerlesen_okul VARCHAR(200)")
    except:
        pass
    try:
        cursor.execute("ALTER TABLE tbl_kullanicilar ADD COLUMN ad_soyad VARCHAR(100)")
    except:
        pass
    
    cursor.execute("SELECT COUNT(*) FROM tbl_kullanicilar")
    if cursor.fetchone()[0] == 0:
        cursor.execute("INSERT INTO tbl_kullanicilar (kullanici_adi, sifre, rol, ad_soyad) VALUES ('admin', '123456', 'Sistem Yöneticisi', 'Sistem Kurucusu')")
    
    cursor.execute('''
        CREATE TABLE IF NOT EXISTS tbl_ogrenciler (
            tc_kimlik VARCHAR(11) PRIMARY KEY,
            ad_soyad VARCHAR(100),
            dogum_tarihi DATE,
            medeni_durum BOOLEAN,
            mevcut_okul_id INTEGER,
            sinif_seviyesi INTEGER,
            sinif_tekrar_sayisi INTEGER,
            lgs_puani FLOAT,
            obp_puani FLOAT,
            ozel_durum_kodu VARCHAR(50),
            kayit_alani VARCHAR(50),
            tercih_1 VARCHAR(200),
            tercih_2 VARCHAR(200),
            tercih_3 VARCHAR(200)
        )
    ''')

    cursor.execute('''
        CREATE TABLE IF NOT EXISTS tbl_nakil_talepleri (
            talep_id INTEGER PRIMARY KEY AUTOINCREMENT,
            tc_kimlik VARCHAR(11),
            talep_turu VARCHAR(50), 
            islem_durumu VARCHAR(20)
        )
    ''')

    cursor.execute('''
        CREATE TABLE IF NOT EXISTS tbl_okullar (
            okul_id INTEGER PRIMARY KEY AUTOINCREMENT,
            okul_adi VARCHAR(200),
            okul_turu VARCHAR(50), 
            kontenjan INTEGER,
            taban_puan FLOAT,
            ozel_egitim_kontenjan INTEGER DEFAULT 10
        )
    ''')
    
    try:
        cursor.execute("ALTER TABLE tbl_okullar ADD COLUMN ozel_egitim_kontenjan INTEGER DEFAULT 10")
    except:
        pass

    conn.commit()
    conn.close()
    
veritabani_kurulum()

def dinamik_egitim_yili_hesapla():
    bugun = datetime.now()
    yil = bugun.year
    ay = bugun.month
    if ay >= 9: baslangic_yili = yil
    else: baslangic_yili = yil - 1
    if baslangic_yili < 2027: baslangic_yili = 2027
    return datetime(baslangic_yili, 9, 1), f"{baslangic_yili}-{baslangic_yili + 1} Eğitim Öğretim Yılı"

@app.route('/')
def ana_panel():
    if not session.get('oturum_acik'): return redirect(url_for('login'))
    referans, egitim_yili_metni = dinamik_egitim_yili_hesapla()
    
    conn = sqlite3.connect(DB_YOLU)
    cursor = conn.cursor()
    cursor.execute("SELECT COUNT(*) FROM tbl_nakil_talepleri WHERE talep_turu='Örgün Nakil Talebi' AND islem_durumu='Beklemede'")
    orgun_nakil = cursor.fetchone()[0]
    
    cursor.execute("SELECT COUNT(*) FROM tbl_nakil_talepleri WHERE talep_turu='Açık Lise Geçiş Talebi' AND islem_durumu='Beklemede'")
    acik_lise = cursor.fetchone()[0]
    
    cursor.execute("SELECT COUNT(*) FROM tbl_nakil_talepleri WHERE talep_turu='MESEM Geçiş Talebi' AND islem_durumu='Beklemede'")
    mesem = cursor.fetchone()[0]
    conn.close()
    
    istatistikler = {
        'orgun_nakil': orgun_nakil,
        'acik_lise': acik_lise,
        'mesem': mesem,
        'guncel_donem': egitim_yili_metni
    }
    return render_template('dashboard.html', veriler=istatistikler)

@app.route('/tutanak_indir')
@tam_yetki_gerekli
def tutanak_indir():
    if not session.get('oturum_acik'): return redirect(url_for('login'))
    
    try:
        from flask import send_file
        import io
        import pandas as pd
        from datetime import datetime
        import sqlite3
        
        conn = sqlite3.connect(DB_YOLU)
        cursor = conn.cursor()
        
        cursor.execute('''
            SELECT o.tc_kimlik, o.ad_soyad, o.dogum_tarihi, t.yerlesen_okul, t.talep_turu
            FROM tbl_ogrenciler o
            INNER JOIN tbl_nakil_talepleri t ON o.tc_kimlik = t.tc_kimlik
            WHERE t.islem_durumu = 'Onaylandı'
        ''')
        onaylananlar = cursor.fetchall()
        
        cursor.execute("SELECT ad_soyad, rol FROM tbl_kullanicilar WHERE rol='Şube Müdürü-Komisyon Başkanı'")
        baskan_verisi = cursor.fetchone()
        baskan_ad = baskan_verisi[0] if baskan_verisi else "..............................."
        
        cursor.execute("SELECT ad_soyad, rol FROM tbl_kullanicilar WHERE rol='Komisyon Üyesi'")
        uyeler = cursor.fetchall()
        conn.close()

        rapor_verisi = []
        sira = 1
        for ogr in onaylananlar:
            okul_adi = ogr[3]
            talep = ogr[4]
            
            if not okul_adi:
                if talep == 'MESEM Geçiş Talebi':
                    okul_adi = 'MESLEKİ EĞİTİM MERKEZİ (MESEM)'
                elif talep == 'Açık Lise Geçiş Talebi':
                    okul_adi = 'AÇIK ÖĞRETİM LİSESİ'
                else:
                    okul_adi = ''
            
            dogum_tarihi_ham = str(ogr[2]).strip() if ogr[2] else ""
            dogum_tarihi_formatli = dogum_tarihi_ham
            if dogum_tarihi_ham and len(dogum_tarihi_ham) >= 10:
                try:
                    dt_obj = datetime.strptime(dogum_tarihi_ham[:10], '%Y-%m-%d')
                    dogum_tarihi_formatli = dt_obj.strftime('%d/%m/%Y')
                except ValueError:
                    pass

            rapor_verisi.append({
                "Sıra No": sira,
                "Adı Soyadı": ogr[1],
                "Doğum Tarihi": dogum_tarihi_formatli,
                "T.C. Kimlik No": ogr[0],
                "Kayıtlı Olduğu Okul": "",
                "Yerleştirildiği Okul": okul_adi,
                "Ort. Öğrt. Yön. İlgili Maddesi/Nakil Sebebi": "",
                "Yerleştirilme Durumu": "UYGUNDUR"
            })
            sira += 1

        df = pd.DataFrame(rapor_verisi)
        output = io.BytesIO()
        
        with pd.ExcelWriter(output, engine='openpyxl') as writer:
            df.to_excel(writer, index=False, sheet_name='Karar Tutanağı', startrow=6)
            worksheet = writer.sheets['Karar Tutanağı']
            
            from openpyxl.styles import Border, Side, Alignment, Font, PatternFill
            ince_cizgi = Side(style='thin')
            kenarlik = Border(left=ince_cizgi, right=ince_cizgi, top=ince_cizgi, bottom=ince_cizgi)
            
            sola_dayali = Alignment(horizontal='left', vertical='center', wrap_text=True)
            ortali = Alignment(horizontal='center', vertical='center', wrap_text=True)
            gri_arkaplan = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid")
            kalin_font = Font(bold=True)
            
            worksheet.merge_cells('A1:H1')
            worksheet['A1'] = "EYÜPSULTAN İLÇE MİLLÎ EĞİTİM MÜDÜRLÜĞÜ"
            worksheet['A1'].font = Font(bold=True, size=14)
            worksheet['A1'].alignment = ortali
            
            worksheet.merge_cells('A2:H2')
            worksheet['A2'] = "İLÇE ÖĞRENCİ YERLEŞTİRME VE NAKİL KOMİSYONU KARAR TUTANAĞI"
            worksheet['A2'].font = Font(bold=True, size=12)
            worksheet['A2'].alignment = ortali
            
            worksheet.merge_cells('A3:H3')

            worksheet.merge_cells('A4:B4')
            worksheet['A4'] = "KARAR TARİHİ :"
            worksheet['A4'].font = kalin_font
            worksheet['A4'].alignment = sola_dayali
            worksheet.merge_cells('C4:H4') 
            
            worksheet.merge_cells('A5:B5')
            worksheet['A5'] = "KARAR NO :"
            worksheet['A5'].font = kalin_font
            worksheet['A5'].alignment = sola_dayali
            worksheet.merge_cells('C5:H5') 
            
            worksheet.merge_cells('A6:H6')

            for row_cells in worksheet.iter_rows(min_row=1, max_row=6, min_col=1, max_col=8):
                for cell in row_cells:
                    cell.border = kenarlik

            for col_idx, cell in enumerate(worksheet[7], 1):
                cell.font = kalin_font
                cell.alignment = ortali
                cell.border = kenarlik
                cell.fill = gri_arkaplan

            for row in worksheet.iter_rows(min_row=8, max_row=len(df) + 7, min_col=1, max_col=8):
                for cell in row:
                    cell.border = kenarlik
                    cell.alignment = sola_dayali
                    
            for i in range(1, len(df) + 8):
                worksheet.row_dimensions[i].height = 35

            worksheet.column_dimensions['A'].width = 8
            worksheet.column_dimensions['B'].width = 20
            worksheet.column_dimensions['C'].width = 12
            worksheet.column_dimensions['D'].width = 16
            worksheet.column_dimensions['E'].width = 25
            worksheet.column_dimensions['F'].width = 45
            worksheet.column_dimensions['G'].width = 26
            worksheet.column_dimensions['H'].width = 18

            imza_satiri = len(df) + 10 
            
            worksheet.cell(row=imza_satiri, column=2, value="ŞUBE MÜDÜRÜ - KOMİSYON BAŞKANI").font = kalin_font
            worksheet.cell(row=imza_satiri, column=2).alignment = ortali
            worksheet.cell(row=imza_satiri + 1, column=2, value=baskan_ad).alignment = ortali
            
            sutun = 4
            for uye in uyeler:
                if sutun > 8: break
                worksheet.cell(row=imza_satiri, column=sutun, value=str(uye[1]).upper()).font = kalin_font
                worksheet.cell(row=imza_satiri, column=sutun).alignment = ortali
                worksheet.cell(row=imza_satiri + 1, column=sutun, value=str(uye[0])).alignment = ortali
                sutun += 2

        output.seek(0)
        tarih = datetime.now().strftime("%d-%m-%Y")
        dosya_adi = f"Eyupsultan_Nakil_Tutanagi_{tarih}.xlsx"
        
        try:
            return send_file(output, download_name=dosya_adi, as_attachment=True)
        except TypeError:
            return send_file(output, attachment_filename=dosya_adi, as_attachment=True)

    except Exception as e:
        return f"<div style='padding:20px;font-family:sans-serif;'><h3>Excel hatası:</h3><p>{str(e)}</p></div>"

@app.route('/kullanicilar')
@tam_yetki_gerekli
def kullanicilar():
    if not session.get('oturum_acik'): return redirect(url_for('login'))
    conn = sqlite3.connect(DB_YOLU)
    conn.row_factory = sqlite3.Row
    cursor = conn.cursor()
    cursor.execute("SELECT * FROM tbl_kullanicilar")
    liste = cursor.fetchall()
    conn.close()
    return render_template('kullanicilar.html', kullanicilar=liste)

@app.route('/kullanici_ekle', methods=['POST'])
@tam_yetki_gerekli
def kullanici_ekle():
    if not session.get('oturum_acik'): return redirect(url_for('login'))
    tc_kimlik = request.form.get('tc_kimlik')
    ad_soyad = request.form.get('ad_soyad')
    rol = request.form.get('komisyon_gorevi')
    
    sifre = guvenli_sifre_olustur()
    
    conn = sqlite3.connect(DB_YOLU)
    cursor = conn.cursor()
    try:
        cursor.execute("INSERT INTO tbl_kullanicilar (kullanici_adi, sifre, rol, ad_soyad) VALUES (?, ?, ?, ?)", (tc_kimlik, sifre, rol, ad_soyad))
        conn.commit()
        flash(f"Kullanıcı eklendi! Sisteme Giriş Şifresi: {sifre} (Lütfen kullanıcıya iletiniz)", "success")
    except sqlite3.IntegrityError:
        flash("Bu TC Kimlik numarasına sahip bir kullanıcı zaten mevcut!", "danger")
    finally:
        conn.close()
        
    return redirect(url_for('kullanicilar'))

@app.route('/kullanici_sil/<kullanici_adi>')
@tam_yetki_gerekli
def kullanici_sil(kullanici_adi):
    if not session.get('oturum_acik'): return redirect(url_for('login'))
            
    conn = sqlite3.connect(DB_YOLU)
    cursor = conn.cursor()
    cursor.execute("DELETE FROM tbl_kullanicilar WHERE kullanici_adi=?", (kullanici_adi,))
    conn.commit()
    conn.close()
    flash(f"Kullanıcı ({kullanici_adi}) başarıyla silindi ve yetkileri alındı.", "info")
    return redirect(url_for('kullanicilar'))

@app.route('/ogrenci_ekle', methods=['GET', 'POST'])
def ogrenci_ekle():
    if not session.get('oturum_acik'): return redirect(url_for('login'))
    if request.method == 'POST':
        tc = request.form.get('tc_kimlik')
        ad = request.form.get('ad_soyad')
        dogum = request.form.get('dogum_tarihi')
        ozel = request.form.get('ozel_durum')
        
        kayit = 'Kayit_Alaninda' 
        talep = request.form.get('talep_turu')
        
        t1 = request.form.get('tercih_1', '').strip()
        t2 = request.form.get('tercih_2', '').strip()
        t3 = request.form.get('tercih_3', '').strip()
        lgs = float(request.form.get('lgs_puani') or 0.0)
        obp = float(request.form.get('obp_puani') or 0.0)
        
        conn = sqlite3.connect(DB_YOLU)
        cursor = conn.cursor()
        
        cursor.execute('''
            INSERT OR REPLACE INTO tbl_ogrenciler 
            (tc_kimlik, ad_soyad, dogum_tarihi, lgs_puani, obp_puani, ozel_durum_kodu, kayit_alani, tercih_1, tercih_2, tercih_3)
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
        ''', (tc, ad, dogum, lgs, obp, ozel, kayit, t1, t2, t3))
        
        if talep and talep != 'Yeni Kayıt':
            cursor.execute("SELECT COUNT(*) FROM tbl_nakil_talepleri WHERE tc_kimlik=? AND talep_turu=? AND islem_durumu='Beklemede'", (tc, talep))
            if cursor.fetchone()[0] == 0:
                cursor.execute("INSERT INTO tbl_nakil_talepleri (tc_kimlik, talep_turu, islem_durumu) VALUES (?, ?, 'Beklemede')", (tc, talep))
                
        conn.commit()
        conn.close()
        return redirect(url_for('ogrenci_havuzu'))
    return render_template('ogrenci_ekle.html')

@app.route('/excel_yukle', methods=['POST'])
def excel_yukle():
    if not session.get('oturum_acik'): return redirect(url_for('login'))
    if 'excel_dosyasi' not in request.files: return redirect(url_for('ogrenci_ekle'))
        
    file = request.files['excel_dosyasi']
    if file and (file.filename.endswith('.xlsx') or file.filename.endswith('.xls')):
        df = pd.read_excel(file)
        
        df.columns = df.columns.astype(str).str.strip()
        
        def sutun_bul(aranan_kelimeler):
            for col in df.columns:
                temiz_col = col.lower().replace(' ', '').replace(':', '').replace('.', '')
                for aranan in aranan_kelimeler:
                    if aranan in temiz_col: return col
            return None
            
        col_tc = sutun_bul(['tc', 'kimlik'])
        col_ad = sutun_bul(['adsoyad', 'adi'])
        col_dogum = sutun_bul(['dogum', 'doğum'])
        col_talep = sutun_bul(['talep'])
        col_lgs = sutun_bul(['lgs'])
        col_obp = sutun_bul(['obp'])
        col_ozel = sutun_bul(['ozel', 'özel'])
        col_t1 = sutun_bul(['tercih1'])
        col_t2 = sutun_bul(['tercih2'])
        col_t3 = sutun_bul(['tercih3'])

        conn = sqlite3.connect(DB_YOLU)
        cursor = conn.cursor()
        
        for index, row in df.iterrows():
            if not col_tc or pd.isna(row[col_tc]): continue
            
            tc = str(row[col_tc]).strip()
            if tc.endswith('.0'): tc = tc[:-2]
            
            if len(tc) < 11 and tc.isdigit(): tc = tc.zfill(11)
            if not tc or tc == 'nan': continue
            
            ad = str(row[col_ad]).strip() if col_ad and pd.notna(row[col_ad]) else ""
            
            dogum_ham = row[col_dogum] if col_dogum else ""
            dogum = ""
            if pd.notna(dogum_ham):
                if isinstance(dogum_ham, datetime):
                    dogum = dogum_ham.strftime('%Y-%m-%d')
                else:
                    try:
                        dt = pd.to_datetime(dogum_ham, dayfirst=True)
                        dogum = dt.strftime('%Y-%m-%d')
                    except:
                        dogum = str(dogum_ham)[:10]
            
            kayit = 'Kayit_Alaninda'
            talep = str(row[col_talep]).strip() if col_talep and pd.notna(row[col_talep]) else 'Yeni Kayıt'
            
            t1 = str(row[col_t1]).strip() if col_t1 and pd.notna(row[col_t1]) else ''
            t2 = str(row[col_t2]).strip() if col_t2 and pd.notna(row[col_t2]) else ''
            t3 = str(row[col_t3]).strip() if col_t3 and pd.notna(row[col_t3]) else ''
            
            def puan_duzelt(deger, max_val):
                if pd.isna(deger): return 0.0
                s = str(deger).strip().replace(',', '.')
                try:
                    f = float(s)
                    while f > max_val: f = f / 10.0
                    return round(f, 4)
                except: return 0.0
            
            lgs = puan_duzelt(row[col_lgs] if col_lgs else 0.0, 500)
            obp = puan_duzelt(row[col_obp] if col_obp else 0.0, 100)
            
            ozel_ham = str(row[col_ozel]).strip() if col_ozel and pd.notna(row[col_ozel]) else 'Yok'
            ozel_kodu = 'Yok'
            
            ozel_norm = ozel_ham.lower().replace('_', ' ').replace('ğ', 'g').replace('ı', 'i').replace('ö', 'o').replace('ü', 'u').replace('ş', 's').replace('ç', 'c')

            if 'ozel egitim' in ozel_norm or 'ozelegitim' in ozel_norm: 
                ozel_kodu = 'Ozel_Egitim'
            elif 'saglik' in ozel_norm: 
                ozel_kodu = 'Saglik_Raporu'
            elif 'anne' in ozel_norm or 'baba' in ozel_norm or 'vefat' in ozel_norm: 
                ozel_kodu = 'Vefat_Anne_Baba'
            elif 'milli' in ozel_norm: 
                ozel_kodu = 'Milli_Sporcu'
            elif 'tutuk' in ozel_norm or 'hukum' in ozel_norm: 
                ozel_kodu = 'Tutuklu_Hukumlu'
            elif 'gazi' in ozel_norm or 'sehit' in ozel_norm: 
                ozel_kodu = 'Gazi_Sehit'
            elif 'korunma' in ozel_norm: 
                ozel_kodu = 'Korunma_Karari'
            elif 'yurt' in ozel_norm: 
                ozel_kodu = 'Yurt_Disi'
            elif ozel_ham in ['Saglik_Raporu', 'Vefat_Anne_Baba', 'Milli_Sporcu', 'Tutuklu_Hukumlu', 'Gazi_Sehit', 'Korunma_Karari', 'Ozel_Egitim', 'Yurt_Disi']:
                ozel_kodu = ozel_ham

            if talep == 'nan' or not talep: talep = 'Yeni Kayıt'
            
            try:
                cursor.execute('''
                    INSERT OR REPLACE INTO tbl_ogrenciler 
                    (tc_kimlik, ad_soyad, dogum_tarihi, lgs_puani, obp_puani, ozel_durum_kodu, kayit_alani, tercih_1, tercih_2, tercih_3)
                    VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                ''', (tc, ad, dogum, lgs, obp, ozel_kodu, kayit, t1, t2, t3))
                
                if talep != 'Yeni Kayıt':
                    cursor.execute("SELECT COUNT(*) FROM tbl_nakil_talepleri WHERE tc_kimlik=? AND talep_turu=? AND islem_durumu='Beklemede'", (tc, talep))
                    if cursor.fetchone()[0] == 0:
                        cursor.execute("INSERT INTO tbl_nakil_talepleri (tc_kimlik, talep_turu, islem_durumu) VALUES (?, ?, 'Beklemede')", (tc, talep))
            except: pass 
                
        conn.commit()
        conn.close()
    return redirect(url_for('ogrenci_havuzu'))

@app.route('/ogrenci_sil/<tc_kimlik>')
@tam_yetki_gerekli
def ogrenci_sil(tc_kimlik):
    if not session.get('oturum_acik'): return redirect(url_for('login'))
    conn = sqlite3.connect(DB_YOLU)
    cursor = conn.cursor()
    cursor.execute("DELETE FROM tbl_ogrenciler WHERE tc_kimlik=?", (str(tc_kimlik).strip(),))
    cursor.execute("DELETE FROM tbl_nakil_talepleri WHERE tc_kimlik=?", (str(tc_kimlik).strip(),))
    conn.commit()
    conn.close()
    return redirect(url_for('ogrenci_havuzu'))

@app.route('/tum_ogrencileri_sil')
@tam_yetki_gerekli
def tum_ogrencileri_sil():
    if not session.get('oturum_acik'): return redirect(url_for('login'))
    conn = sqlite3.connect(DB_YOLU)
    cursor = conn.cursor()
    cursor.execute("DELETE FROM tbl_ogrenciler")
    cursor.execute("DELETE FROM tbl_nakil_talepleri")
    conn.commit()
    conn.close()
    return redirect(url_for('ogrenci_havuzu'))

@app.route('/havuz')
def ogrenci_havuzu():
    conn = sqlite3.connect(DB_YOLU)
    conn.row_factory = sqlite3.Row 
    cursor = conn.cursor()
    cursor.execute("SELECT * FROM tbl_ogrenciler ORDER BY obp_puani DESC")
    ogrenciler = cursor.fetchall()
    conn.close()
    return render_template('havuz.html', ogrenciler=ogrenciler)

@app.route('/okul_ekle', methods=['GET', 'POST'])
@tam_yetki_gerekli
def okul_ekle():
    if not session.get('oturum_acik'): return redirect(url_for('login'))
    if request.method == 'POST':
        adi = request.form.get('okul_adi')
        turu = request.form.get('okul_turu')
        kontenjan = int(request.form.get('kontenjan', 0))
        taban_puan = float(request.form.get('taban_puan', 0))
        ozel_kontenjan = int(request.form.get('ozel_egitim_kontenjan') or 10)
        
        conn = sqlite3.connect(DB_YOLU)
        cursor = conn.cursor()
        cursor.execute('INSERT INTO tbl_okullar (okul_adi, okul_turu, kontenjan, taban_puan, ozel_egitim_kontenjan) VALUES (?, ?, ?, ?, ?)', (adi, turu, kontenjan, taban_puan, ozel_kontenjan))
        conn.commit()
        conn.close()
        return redirect(url_for('okullar_listesi'))
    return render_template('okul_ekle.html')

@app.route('/okul_sil/<okul_id>')
@tam_yetki_gerekli
def okul_sil(okul_id):
    if not session.get('oturum_acik'): return redirect(url_for('login'))
    conn = sqlite3.connect(DB_YOLU)
    cursor = conn.cursor()
    cursor.execute("DELETE FROM tbl_okullar WHERE okul_id=?", (int(okul_id),))
    conn.commit()
    conn.close()
    return redirect(url_for('okullar_listesi'))

@app.route('/okul_duzenle/<int:okul_id>', methods=['GET', 'POST'])
@tam_yetki_gerekli
def okul_duzenle(okul_id):
    if not session.get('oturum_acik'): return redirect(url_for('login'))
    conn = sqlite3.connect(DB_YOLU)
    conn.row_factory = sqlite3.Row
    cursor = conn.cursor()
    
    if request.method == 'POST':
        adi = request.form.get('okul_adi')
        turu = request.form.get('okul_turu')
        kontenjan = int(request.form.get('kontenjan', 0))
        taban_puan = float(request.form.get('taban_puan', 0))
        ozel_kontenjan = int(request.form.get('ozel_egitim_kontenjan') or 10)
        
        cursor.execute('''
            UPDATE tbl_okullar 
            SET okul_adi=?, okul_turu=?, kontenjan=?, taban_puan=?, ozel_egitim_kontenjan=? 
            WHERE okul_id=?
        ''', (adi, turu, kontenjan, taban_puan, ozel_kontenjan, okul_id))
        
        conn.commit()
        conn.close()
        return redirect(url_for('okullar_listesi'))
        
    cursor.execute("SELECT * FROM tbl_okullar WHERE okul_id=?", (okul_id,))
    okul = cursor.fetchone()
    conn.close()
    return render_template('okul_duzenle.html', okul=okul)

@app.route('/okullar')
def okullar_listesi():
    conn = sqlite3.connect(DB_YOLU)
    conn.row_factory = sqlite3.Row
    cursor = conn.cursor()
    cursor.execute("SELECT * FROM tbl_okullar")
    okullar = cursor.fetchall()
    conn.close()
    return render_template('okullar.html', okullar=okullar)

@app.route('/eslestirme')
def eslestirme_motoru():
    if not session.get('oturum_acik'): return redirect(url_for('login'))
    conn = sqlite3.connect(DB_YOLU)
    conn.row_factory = sqlite3.Row
    cursor = conn.cursor()
    
    cursor.execute("SELECT * FROM tbl_okullar")
    okullar_db = cursor.fetchall()
    okullar = {}
    for row in okullar_db:
        o_kontenjan = row['ozel_egitim_kontenjan'] if 'ozel_egitim_kontenjan' in row.keys() and row['ozel_egitim_kontenjan'] is not None else 10
        
        okullar[row['okul_adi'].strip()] = {
            'turu': str(row['okul_turu']).lower(),
            'kontenjan': int(row['kontenjan'] or 0),
            'taban': float(row['taban_puan'] or 0.0),
            'yerlesenler': [],
            'standart_yerlesen_sayisi': 0, 
            'ozel_egitim_kalan': int(o_kontenjan),
            'orijinal_isim': row['okul_adi']
        }
        
    cursor.execute('''
        SELECT o.* 
        FROM tbl_ogrenciler o
        LEFT JOIN tbl_nakil_talepleri t ON o.tc_kimlik = t.tc_kimlik
        WHERE t.talep_turu IS NULL OR t.talep_turu = 'Örgün Nakil Talebi'
    ''')
    tum_ogrenciler = [dict(row) for row in cursor.fetchall()]
    
    cursor.execute('''
        SELECT o.*, t.talep_turu, t.islem_durumu 
        FROM tbl_ogrenciler o
        INNER JOIN tbl_nakil_talepleri t ON o.tc_kimlik = t.tc_kimlik
        WHERE t.talep_turu IN ('Açık Lise Geçiş Talebi', 'MESEM Geçiş Talebi')
    ''')
    ozel_komisyon_ogrencileri = [dict(row) for row in cursor.fetchall()]
    
    conn.close()
    
    yerlesim_sonuclari = {}
    yerlesen_tc = set()
    
    def is_lgs(turu): return any(k in turu for k in ['merkezi', 'lgs', 'sınav', 'sinav'])
    
    def ozel_mi(durum): 
        durum = str(durum).strip()
        return 'Yok' not in durum and durum not in ['None', '']
        
    def is_ozel_egitim(durum):
        d = str(durum).strip().lower().replace('_', ' ').replace('ğ', 'g').replace('ı', 'i').replace('ö', 'o').replace('ü', 'u').replace('ş', 's').replace('ç', 'c')
        return 'ozel egitim' in d or 'ozelegitim' in d

    def kayit_puani(alan): return 3 if alan == 'Kayit_Alaninda' else 2 if alan == 'Komsu_Kayit_Alani' else 1

    lgs_adaylari = [o for o in tum_ogrenciler if float(o.get('lgs_puani') or 0.0) > 0]
    lgs_adaylari.sort(key=lambda x: (float(x.get('lgs_puani') or 0.0), x.get('dogum_tarihi', '')), reverse=True)
    
    for ogrenci in lgs_adaylari:
        if ogrenci['tc_kimlik'] in yerlesen_tc: continue
        tercihler = [str(ogrenci.get(f'tercih_{i}') or '').strip() for i in range(1, 4)]
        for tercih in tercihler:
            if tercih in okullar and is_lgs(okullar[tercih]['turu']):
                hedef = okullar[tercih]
                is_muaf = ozel_mi(ogrenci.get('ozel_durum_kodu'))
                is_oe = is_ozel_egitim(ogrenci.get('ozel_durum_kodu'))
                
                if is_oe:
                    if hedef['ozel_egitim_kalan'] <= 0: continue
                    hedef['yerlesenler'].append(ogrenci)
                    yerlesen_tc.add(ogrenci['tc_kimlik'])
                    hedef['ozel_egitim_kalan'] -= 1
                    break
                elif is_muaf and not is_oe:
                    hedef['yerlesenler'].append(ogrenci)
                    yerlesen_tc.add(ogrenci['tc_kimlik'])
                    hedef['standart_yerlesen_sayisi'] += 1
                    break
                else:
                    if hedef['standart_yerlesen_sayisi'] < hedef['kontenjan']:
                        if float(ogrenci.get('lgs_puani') or 0.0) >= hedef['taban']:
                            hedef['yerlesenler'].append(ogrenci)
                            yerlesen_tc.add(ogrenci['tc_kimlik'])
                            hedef['standart_yerlesen_sayisi'] += 1
                            break 

    obp_adaylari = [o for o in tum_ogrenciler if o['tc_kimlik'] not in yerlesen_tc]
    obp_adaylari.sort(key=lambda x: (kayit_puani(x.get('kayit_alani', 'Diger')), float(x.get('obp_puani') or 0.0), x.get('dogum_tarihi', '')), reverse=True)
    
    for ogrenci in obp_adaylari:
        tercihler = [str(ogrenci.get(f'tercih_{i}') or '').strip() for i in range(1, 4)]
        for tercih in tercihler:
            if tercih in okullar and not is_lgs(okullar[tercih]['turu']):
                hedef = okullar[tercih]
                is_muaf = ozel_mi(ogrenci.get('ozel_durum_kodu'))
                is_oe = is_ozel_egitim(ogrenci.get('ozel_durum_kodu'))
                
                if is_oe:
                    if hedef['ozel_egitim_kalan'] <= 0: continue
                    hedef['yerlesenler'].append(ogrenci)
                    yerlesen_tc.add(ogrenci['tc_kimlik'])
                    hedef['ozel_egitim_kalan'] -= 1
                    break
                elif is_muaf and not is_oe:
                    hedef['yerlesenler'].append(ogrenci)
                    yerlesen_tc.add(ogrenci['tc_kimlik'])
                    hedef['standart_yerlesen_sayisi'] += 1
                    break
                else:
                    if hedef['standart_yerlesen_sayisi'] < hedef['kontenjan']:
                        hedef['yerlesenler'].append(ogrenci)
                        yerlesen_tc.add(ogrenci['tc_kimlik'])
                        hedef['standart_yerlesen_sayisi'] += 1
                        break

    kalan_adaylar = [o for o in tum_ogrenciler if o['tc_kimlik'] not in yerlesen_tc]
    kalan_adaylar.sort(key=lambda x: (kayit_puani(x.get('kayit_alani', 'Diger')), float(x.get('obp_puani') or 0.0), x.get('dogum_tarihi', '')), reverse=True)
    
    bos_okullar = [okul for okul in okullar.values() if not is_lgs(okul['turu']) and (okul['standart_yerlesen_sayisi'] < okul['kontenjan'] or okul['ozel_egitim_kalan'] > 0)]
    
    def okul_onceligi(okul_adi):
        ad_lower = okul_adi.lower()
        if 'meslek' in ad_lower or 'mesleki' in ad_lower or 'teknik' in ad_lower: return 2
        if 'imam' in ad_lower or 'hatip' in ad_lower: return 2
        return 1

    bos_okullar.sort(key=lambda x: (okul_onceligi(x['orijinal_isim']), x['kontenjan'] - x['standart_yerlesen_sayisi']), reverse=True)

    for ogrenci in kalan_adaylar:
        is_oe = is_ozel_egitim(ogrenci.get('ozel_durum_kodu'))
        for hedef_okul in bos_okullar:
            if is_oe:
                if hedef_okul['ozel_egitim_kalan'] > 0:
                    hedef_okul['yerlesenler'].append(ogrenci)
                    yerlesen_tc.add(ogrenci['tc_kimlik'])
                    hedef_okul['ozel_egitim_kalan'] -= 1
                    break
            else:
                if hedef_okul['standart_yerlesen_sayisi'] < hedef_okul['kontenjan']:
                    hedef_okul['yerlesenler'].append(ogrenci)
                    yerlesen_tc.add(ogrenci['tc_kimlik'])
                    hedef_okul['standart_yerlesen_sayisi'] += 1
                    break 

    for ad, okul in okullar.items():
        yerlesim_sonuclari[okul['orijinal_isim']] = {
            'turu': 'LGS' if is_lgs(okul['turu']) else 'OBP',
            'ogrenciler': okul['yerlesenler'],
            'doluluk': f"{okul['standart_yerlesen_sayisi']} / {okul['kontenjan']} (Özel Eğitim Kontenjanı: {okul['ozel_egitim_kalan']} Kalan)"
        }
        
    gercekten_acikta_kalanlar = [o for o in tum_ogrenciler if o['tc_kimlik'] not in yerlesen_tc]
    
    if yerlesen_tc:
        db_conn = sqlite3.connect(DB_YOLU)
        db_cursor = db_conn.cursor()
        
        for okul_adi, okul_bilgisi in yerlesim_sonuclari.items():
            for kazanan_ogrenci in okul_bilgisi['ogrenciler']:
                db_cursor.execute("""
                    UPDATE tbl_nakil_talepleri 
                    SET islem_durumu='Onaylandı', yerlesen_okul=?
                    WHERE tc_kimlik=?
                """, (okul_adi, kazanan_ogrenci['tc_kimlik']))
        db_conn.commit()
        db_conn.close()
    return render_template('eslestirme.html', 
                           sonuclar=yerlesim_sonuclari, 
                           aciktakiler=gercekten_acikta_kalanlar, 
                           komisyon_bekleyenler=ozel_komisyon_ogrencileri)

@app.route('/komisyon')
def komisyon_paneli():
    if not session.get('oturum_acik'): return redirect(url_for('login'))
    
    conn = sqlite3.connect(DB_YOLU)
    conn.row_factory = sqlite3.Row
    cursor = conn.cursor()
    cursor.execute('''
        SELECT o.tc_kimlik, o.ad_soyad, t.talep_id, t.talep_turu, t.islem_durumu 
        FROM tbl_ogrenciler o
        INNER JOIN tbl_nakil_talepleri t ON o.tc_kimlik = t.tc_kimlik
        WHERE t.talep_turu IN ('Açık Lise Geçiş Talebi', 'MESEM Geçiş Talebi')
        ORDER BY t.islem_durumu ASC, o.ad_soyad ASC
    ''')
    talepler = cursor.fetchall()
    conn.close()
    
    return render_template('komisyon.html', talepler=talepler)

@app.route('/talep_guncelle/<int:talep_id>/<durum>')
@tam_yetki_gerekli
def talep_guncelle(talep_id, durum):
    if not session.get('oturum_acik'): return redirect(url_for('login'))
    if durum not in ['Onaylandı', 'Reddedildi', 'Beklemede']: return redirect(url_for('komisyon_paneli'))
    
    conn = sqlite3.connect(DB_YOLU)
    cursor = conn.cursor()
    cursor.execute("UPDATE tbl_nakil_talepleri SET islem_durumu=? WHERE talep_id=?", (durum, talep_id))
    conn.commit()
    conn.close()
    
    return redirect(url_for('komisyon_paneli'))

@app.route('/login', methods=['GET', 'POST'])
def login():
    if request.method == 'POST':
        kullanici = request.form.get('kullanici_adi')
        sifre = request.form.get('sifre')
        conn = sqlite3.connect(DB_YOLU)
        cursor = conn.cursor()
        cursor.execute("SELECT * FROM tbl_kullanicilar WHERE kullanici_adi=? AND sifre=?", (kullanici, sifre))
        kullanici_verisi = cursor.fetchone()
        conn.close()
        
        if kullanici_verisi:
            session['oturum_acik'] = True
            session['kullanici_rolu'] = kullanici_verisi[2] 
            return redirect(url_for('ana_panel'))
        else:
            return render_template('login.html', hata="Hatalı TC veya şifre!")
    return render_template('login.html')

@app.route('/logout')
def logout():
    session.clear() 
    return redirect(url_for('login'))

if __name__ == '__main__':
    app.run(host='0.0.0.0', port=5000, debug=True)